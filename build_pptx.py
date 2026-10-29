"""
ContractAI - Professional 10-Slide Presentation Generator
Based on actual project implementation inspection.
Uses: python-pptx
"""

import os
from pptx import Presentation
from pptx.util import Inches, Pt, Emu
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN

# ─── COLOR PALETTE ───────────────────────────────────────────────────────────
NAVY        = RGBColor(0x0D, 0x1B, 0x2A)
CHARCOAL    = RGBColor(0x1E, 0x29, 0x3B)
SLATE       = RGBColor(0x33, 0x4A, 0x5C)
WHITE       = RGBColor(0xFF, 0xFF, 0xFF)
OFF_WHITE   = RGBColor(0xF1, 0xF5, 0xF9)
INDIGO      = RGBColor(0x6D, 0x28, 0xD9)
INDIGO_LT   = RGBColor(0x7C, 0x3A, 0xED)
PURPLE_DIM  = RGBColor(0x3B, 0x0F, 0x6E)
GRAY_MED    = RGBColor(0x94, 0xA3, 0xB8)
GRAY_LIGHT  = RGBColor(0xCB, 0xD5, 0xE1)
GREEN_OK    = RGBColor(0x10, 0xB9, 0x81)
AMBER_MED   = RGBColor(0xF5, 0x9E, 0x0B)
RED_HIGH    = RGBColor(0xEF, 0x44, 0x44)
DARK_PURPLE = RGBColor(0x1A, 0x07, 0x40)
TEAL_DARK   = RGBColor(0x06, 0x50, 0x3C)
CRIMSON     = RGBColor(0x5A, 0x1A, 0x1A)

BRAIN_DIR = r"C:\Users\ayush\.gemini\antigravity\brain\d0209986-f3ee-4f1c-aaf1-66c4bed99fbe"
DASHBOARD_IMG  = os.path.join(BRAIN_DIR, "dashboard_page_1790395267643.png")
UPLOAD_IMG     = os.path.join(BRAIN_DIR, "upload_page_1790395287144.png")
AGREEMENTS_IMG = os.path.join(BRAIN_DIR, "agreements_page_1790395321135.png")
COMPARE_IMG    = os.path.join(BRAIN_DIR, "compare_page_1790395336667.png")

OUTPUT_PATH = r"c:\Users\ayush\OneDrive\Desktop\project 2 sem 7\ContractAI_Presentation.pptx"


def new_prs():
    prs = Presentation()
    prs.slide_width  = Inches(13.33)
    prs.slide_height = Inches(7.5)
    return prs


def blank_layout(prs):
    return prs.slide_layouts[6]


def set_slide_bg(slide, color):
    fill = slide.background.fill
    fill.solid()
    fill.fore_color.rgb = color


def add_rect(slide, left, top, width, height, fill=None, line=None, line_w=None):
    sh = slide.shapes.add_shape(1, Inches(left), Inches(top), Inches(width), Inches(height))
    if fill:
        sh.fill.solid()
        sh.fill.fore_color.rgb = fill
    else:
        sh.fill.background()
    if line:
        sh.line.color.rgb = line
        if line_w:
            sh.line.width = Pt(line_w)
    else:
        sh.line.fill.background()
    return sh


def add_tb(slide, text, left, top, width, height,
           size=14, bold=False, italic=False, color=WHITE,
           align=PP_ALIGN.LEFT, font="Calibri", wrap=True):
    tb = slide.shapes.add_textbox(Inches(left), Inches(top), Inches(width), Inches(height))
    tf = tb.text_frame
    tf.word_wrap = wrap
    p = tf.paragraphs[0]
    p.alignment = align
    r = p.add_run()
    r.text = text
    r.font.size = Pt(size)
    r.font.bold = bold
    r.font.italic = italic
    r.font.color.rgb = color
    r.font.name = font
    return tb


def add_mltb(slide, specs, left, top, width, height,
             def_size=13, def_bold=False, def_color=WHITE,
             align=PP_ALIGN.LEFT, font="Calibri"):
    """Multi-paragraph textbox. specs = list of dicts with keys: t, s, b, c, i"""
    tb = slide.shapes.add_textbox(Inches(left), Inches(top), Inches(width), Inches(height))
    tf = tb.text_frame
    tf.word_wrap = True
    for idx, sp in enumerate(specs):
        text  = sp.get('t', '')
        size  = sp.get('s', def_size)
        bold  = sp.get('b', def_bold)
        col   = sp.get('c', def_color)
        ital  = sp.get('i', False)
        if idx == 0:
            p = tf.paragraphs[0]
        else:
            p = tf.add_paragraph()
        p.alignment = align
        r = p.add_run()
        r.text = text
        r.font.size = Pt(size)
        r.font.bold = bold
        r.font.italic = ital
        r.font.color.rgb = col
        r.font.name = font
    return tb


def add_pill_label(slide, text, left, top, width, height,
                   bg=INDIGO, fg=WHITE, size=10):
    add_rect(slide, left, top, width, height, fill=bg)
    add_tb(slide, text, left, top, width, height, size=size, bold=True,
           color=fg, align=PP_ALIGN.CENTER)


def add_slide_header(slide, num, title, sub=None):
    # Number badge
    add_rect(slide, 0.35, 0.22, 0.42, 0.42, fill=INDIGO)
    add_tb(slide, f"{num:02d}", 0.35, 0.22, 0.42, 0.42,
           size=13, bold=True, color=WHITE, align=PP_ALIGN.CENTER)
    # Title
    add_tb(slide, title, 0.90, 0.20, 11.8, 0.52,
           size=32, bold=True, color=WHITE)
    # Subtitle
    if sub:
        add_tb(slide, sub, 0.90, 0.71, 11.5, 0.32,
               size=13, color=GRAY_MED, italic=True)
    # Separator
    add_rect(slide, 0.35, 1.08, 12.63, 0.025, fill=PURPLE_DIM)


def add_footer(slide, text=None):
    if text is None:
        text = "ContractAI  ·  PCTE Institute of Engineering and Technology, Ludhiana  ·  Ayush Yadav  2329304"
    add_tb(slide, text, 0.35, 7.18, 12.63, 0.26,
           size=9, color=GRAY_MED, align=PP_ALIGN.CENTER, italic=True)


# ══════════════════════════════════════════════════════════════════════════════
#  SLIDE 1 — TITLE
# ══════════════════════════════════════════════════════════════════════════════
def slide_1(prs):
    sl = prs.slides.add_slide(blank_layout(prs))
    set_slide_bg(sl, NAVY)

    # Left dark panel
    add_rect(sl, 0, 0, 5.0, 7.5, fill=CHARCOAL)
    # Bottom accent strip
    add_rect(sl, 0, 6.15, 13.33, 1.35, fill=DARK_PURPLE)
    # Indigo vertical glow bar
    add_rect(sl, 4.95, 0, 0.1, 7.5, fill=INDIGO)

    # Left panel content
    add_tb(sl, "PRESENTED BY", 0.35, 0.42, 4.4, 0.28,
           size=9, bold=True, color=INDIGO_LT)
    add_tb(sl, "Ayush Yadav", 0.35, 0.72, 4.4, 0.62,
           size=28, bold=True, color=WHITE)
    add_tb(sl, "Roll No. 2329304", 0.35, 1.35, 4.4, 0.32,
           size=13, color=GRAY_LIGHT)
    add_tb(sl, "B.Tech CSE  ·  Section 4-C", 0.35, 1.65, 4.4, 0.32,
           size=13, color=GRAY_LIGHT)

    add_rect(sl, 0.35, 2.12, 4.2, 0.03, fill=PURPLE_DIM)

    add_tb(sl, "SUBMITTED TO", 0.35, 2.22, 4.4, 0.28,
           size=9, bold=True, color=INDIGO_LT)
    add_tb(sl, "Ms. Mandeep Kaur", 0.35, 2.50, 4.4, 0.42,
           size=16, bold=True, color=WHITE)

    add_rect(sl, 0.35, 3.05, 4.2, 0.03, fill=PURPLE_DIM)

    add_tb(sl, "INSTITUTION", 0.35, 3.15, 4.4, 0.28,
           size=9, bold=True, color=INDIGO_LT)
    add_tb(sl, "PCTE Institute of Engineering\nand Technology, Ludhiana",
           0.35, 3.43, 4.4, 0.72, size=13, color=GRAY_LIGHT)

    add_rect(sl, 0.35, 4.28, 4.2, 0.03, fill=PURPLE_DIM)
    add_tb(sl, "B.Tech CSE  ·  Academic Year 2026–2027", 0.35, 4.38, 4.4, 0.32,
           size=11, color=GRAY_MED, italic=True)

    # Right panel — big brand mark
    add_pill_label(sl, "AI-POWERED LEGALTECH PROJECT", 5.2, 0.35, 7.8, 0.30,
                   bg=PURPLE_DIM, fg=GRAY_LIGHT, size=10)

    add_tb(sl, "CONTRACT", 5.2, 0.82, 7.8, 1.05,
           size=62, bold=True, color=WHITE)
    add_tb(sl, "AI", 5.2, 1.72, 3.5, 0.95,
           size=62, bold=True, color=INDIGO_LT)

    add_rect(sl, 5.2, 2.72, 7.8, 0.06, fill=INDIGO_LT)

    add_tb(sl, "AI-POWERED CONTRACT RISK ANALYZER",
           5.2, 2.88, 7.8, 0.48, size=18, bold=True, color=GRAY_LIGHT)
    add_tb(sl, "Intelligent Contract Analysis Using AI, RAG & Automated Risk Detection",
           5.2, 3.38, 7.8, 0.55, size=13, color=GRAY_MED, italic=True)

    # 3 concept pills
    for i, (lbl, col) in enumerate([
        ("DOCUMENT AI", INDIGO), ("RISK SCORING", RED_HIGH), ("SMART REDLINE", GREEN_OK)
    ]):
        add_pill_label(sl, lbl, 5.2 + i * 2.65, 4.12, 2.52, 0.40,
                       bg=col, fg=WHITE, size=11)

    add_tb(sl, "FastAPI · Python · React/TypeScript · ChromaDB · LangChain · OpenAI · SQLite · ReportLab",
           5.2, 6.38, 7.8, 0.32, size=10, color=GRAY_MED, italic=True)


# ══════════════════════════════════════════════════════════════════════════════
#  SLIDE 2 — THE PROBLEM
# ══════════════════════════════════════════════════════════════════════════════
def slide_2(prs):
    sl = prs.slides.add_slide(blank_layout(prs))
    set_slide_bg(sl, NAVY)
    add_slide_header(sl, 2, "THE PROBLEM", "Why contract review needs automation")
    add_footer(sl)

    pain = [
        ("VOLUME & LENGTH",   "Contracts span dozens of pages with dense legal language"),
        ("MISSED CLAUSES",    "Critical risk clauses can easily be overlooked manually"),
        ("TIME COST",         "Manual legal review is slow and often expensive"),
        ("ACCESSIBILITY",     "Non-legal users struggle to interpret complex legal terms"),
    ]
    for i, (ttl, desc) in enumerate(pain):
        y = 1.35 + i * 1.35
        add_rect(sl, 0.35, y, 6.1, 1.18, fill=CHARCOAL, line=PURPLE_DIM, line_w=0.5)
        add_rect(sl, 0.35, y, 0.09, 1.18, fill=INDIGO_LT)
        add_tb(sl, ttl, 0.60, y + 0.10, 5.6, 0.36,
               size=14, bold=True, color=INDIGO_LT)
        add_tb(sl, desc, 0.60, y + 0.52, 5.6, 0.55,
               size=12, color=GRAY_LIGHT)

    # Right: flow
    add_tb(sl, "TRADITIONAL APPROACH", 7.0, 1.28, 5.95, 0.32,
           size=12, bold=True, color=GRAY_MED, align=PP_ALIGN.CENTER)

    steps = [
        ("📄  LONG CONTRACT",   "Dozens of pages, dense clauses",         RGBColor(0x33, 0x4A, 0x5C)),
        ("👁   MANUAL REVIEW",  "Lawyer reads every clause carefully",     RGBColor(0x28, 0x3A, 0x4E)),
        ("⏱   TIME & COST",    "Hours of billable legal work required",    CRIMSON),
        ("⚠   MISSED RISKS",   "Key risky clauses can be overlooked",      RGBColor(0x7F, 0x1D, 0x1D)),
    ]
    for i, (lbl, sub, bg) in enumerate(steps):
        y = 1.72 + i * 1.15
        add_rect(sl, 6.95, y, 6.0, 0.82, fill=bg)
        add_tb(sl, lbl, 7.08, y + 0.07, 5.75, 0.36,
               size=14, bold=True, color=WHITE, align=PP_ALIGN.CENTER)
        add_tb(sl, sub, 7.08, y + 0.44, 5.75, 0.28,
               size=10, color=GRAY_LIGHT, align=PP_ALIGN.CENTER, italic=True)
        if i < 3:
            add_tb(sl, "↓", 9.7, y + 0.84, 0.55, 0.32,
                   size=16, color=INDIGO_LT, align=PP_ALIGN.CENTER)


# ══════════════════════════════════════════════════════════════════════════════
#  SLIDE 3 — THE SOLUTION
# ══════════════════════════════════════════════════════════════════════════════
def slide_3(prs):
    sl = prs.slides.add_slide(blank_layout(prs))
    set_slide_bg(sl, NAVY)
    add_slide_header(sl, 3, "THE CONTRACTAI SOLUTION",
                     "A browser-based AI contract risk analysis platform")
    add_footer(sl)

    features = [
        ("📤", "Document Ingestion",     "Accepts PDF & DOCX contracts — PyMuPDF + python-docx"),
        ("🔍", "Clause Extraction",      "Identifies 15+ legal clause categories with regex NLP"),
        ("🤖", "AI + RAG Analysis",      "GPT-4o-mini + ChromaDB + Rule Engine (40/40/20 weights)"),
        ("📊", "Risk Score & Redline",   "0–100 composite score with AI-suggested safer clauses"),
    ]
    for i, (icon, title, desc) in enumerate(features):
        y = 1.35 + i * 1.38
        add_rect(sl, 0.35, y, 7.1, 1.2, fill=CHARCOAL, line=PURPLE_DIM, line_w=0.5)
        add_tb(sl, icon, 0.52, y + 0.22, 0.7, 0.58,
               size=24, color=WHITE, align=PP_ALIGN.CENTER)
        add_tb(sl, title, 1.32, y + 0.10, 5.95, 0.38,
               size=15, bold=True, color=INDIGO_LT)
        add_tb(sl, desc, 1.32, y + 0.52, 5.95, 0.55,
               size=12, color=GRAY_LIGHT)

    # Right pipeline
    add_tb(sl, "ANALYSIS PIPELINE", 7.9, 1.28, 5.1, 0.32,
           size=12, bold=True, color=GRAY_MED, align=PP_ALIGN.CENTER)

    pipeline = [
        ("UPLOAD",           GREEN_OK),
        ("TEXT EXTRACTION",  INDIGO_LT),
        ("CLAUSE DETECT",    INDIGO),
        ("RAG / SCORING",    AMBER_MED),
        ("REDLINE",          RED_HIGH),
        ("PDF REPORT",       GREEN_OK),
    ]
    for i, (lbl, col) in enumerate(pipeline):
        y = 1.72 + i * 0.84
        add_rect(sl, 7.82, y, 5.15, 0.60, fill=col)
        add_tb(sl, lbl, 7.82, y + 0.12, 5.15, 0.38,
               size=14, bold=True, color=WHITE, align=PP_ALIGN.CENTER)
        if i < 5:
            add_tb(sl, "▼", 10.22, y + 0.62, 0.55, 0.25,
                   size=12, color=GRAY_LIGHT, align=PP_ALIGN.CENTER)


# ══════════════════════════════════════════════════════════════════════════════
#  SLIDE 4 — SYSTEM WORKFLOW
# ══════════════════════════════════════════════════════════════════════════════
def slide_4(prs):
    sl = prs.slides.add_slide(blank_layout(prs))
    set_slide_bg(sl, NAVY)
    add_slide_header(sl, 4, "SYSTEM WORKFLOW",
                     "End-to-end flow: from document upload to risk dashboard")
    add_footer(sl)

    col1 = [
        ("1", "USER",                    SLATE),
        ("2", "PDF / DOCX UPLOAD",       CHARCOAL),
        ("3", "TEXT EXTRACTION",         CHARCOAL),
        ("4", "CLAUSE EXTRACTION",       RGBColor(0x2A, 0x3A, 0x58)),
        ("5", "METADATA EXTRACTION",     RGBColor(0x2A, 0x3A, 0x58)),
    ]
    col2 = [
        ("6", "ChromaDB / RAG INDEXING", PURPLE_DIM),
        ("7", "AI + RULE ENGINE",        RGBColor(0x4A, 0x15, 0x7A)),
        ("8", "RISK SCORING (0–100)",    CRIMSON),
        ("9", "REDLINE GENERATION",      RGBColor(0x7F, 0x1D, 0x1D)),
        ("10","DASHBOARD + PDF REPORT",  TEAL_DARK),
    ]

    cw, ch = 5.72, 0.72

    for i, (num, lbl, bg) in enumerate(col1):
        y = 1.35 + i * 1.12
        add_rect(sl, 0.40, y, cw, ch, fill=bg)
        add_tb(sl, num, 0.52, y + 0.14, 0.45, 0.38,
               size=13, bold=True, color=INDIGO_LT)
        add_tb(sl, lbl, 1.0, y + 0.14, cw - 0.65, 0.38,
               size=13, bold=True, color=WHITE)
        if i < 4:
            add_tb(sl, "↓", 0.40 + cw / 2 - 0.22, y + ch + 0.02, 0.5, 0.38,
                   size=16, color=INDIGO_LT, align=PP_ALIGN.CENTER)

    # Bridge arrow
    add_tb(sl, "→", 6.18, 3.05, 1.0, 0.65,
           size=32, color=INDIGO_LT, align=PP_ALIGN.CENTER)

    for i, (num, lbl, bg) in enumerate(col2):
        y = 1.35 + i * 1.12
        add_rect(sl, 7.22, y, cw, ch, fill=bg)
        add_tb(sl, num, 7.34, y + 0.14, 0.52, 0.38,
               size=13, bold=True, color=INDIGO_LT)
        add_tb(sl, lbl, 7.90, y + 0.14, cw - 0.72, 0.38,
               size=13, bold=True, color=WHITE)
        if i < 4:
            add_tb(sl, "↓", 7.22 + cw / 2 - 0.22, y + ch + 0.02, 0.5, 0.38,
                   size=16, color=INDIGO_LT, align=PP_ALIGN.CENTER)


# ══════════════════════════════════════════════════════════════════════════════
#  SLIDE 5 — TECHNOLOGY STACK
# ══════════════════════════════════════════════════════════════════════════════
def slide_5(prs):
    sl = prs.slides.add_slide(blank_layout(prs))
    set_slide_bg(sl, NAVY)
    add_slide_header(sl, 5, "TECHNOLOGY STACK",
                     "Technologies verified in the actual implementation")
    add_footer(sl)

    groups = [
        ("BACKEND",             ["Python 3.x", "FastAPI", "Uvicorn", "SQLAlchemy ORM"],               INDIGO),
        ("DOCUMENT PROCESSING", ["PyMuPDF (fitz)", "python-docx", "PDF & DOCX Parsing"],              RGBColor(0x0E, 0x5E, 0x9A)),
        ("AI / LLM",            ["OpenAI GPT-4o-mini", "LangChain", "NLP Rule Matching"],             RGBColor(0x6D, 0x28, 0xD9)),
        ("RAG / VECTOR SEARCH", ["ChromaDB (Persistent)", "Vector Embeddings", "Semantic Search"],    PURPLE_DIM),
        ("FRONTEND",            ["React 19 + TypeScript", "Vite + TailwindCSS", "Recharts", "lucide-react"], TEAL_DARK),
        ("DATABASE & REPORTS",  ["SQLite via SQLAlchemy", "ReportLab PDF Generation"],                CRIMSON),
    ]

    positions = [
        (0.35, 1.32), (4.60, 1.32), (8.85, 1.32),
        (0.35, 4.12), (4.60, 4.12), (8.85, 4.12),
    ]

    cw, ch = 4.05, 2.60

    for (label, techs, accent), (x, y) in zip(groups, positions):
        add_rect(sl, x, y, cw, ch, fill=CHARCOAL, line=PURPLE_DIM, line_w=0.4)
        # Top accent
        add_rect(sl, x, y, cw, 0.10, fill=accent)
        add_tb(sl, label, x + 0.14, y + 0.18, cw - 0.22, 0.34,
               size=12, bold=True, color=accent)
        for j, tech in enumerate(techs[:4]):
            add_tb(sl, f"  •  {tech}", x + 0.12, y + 0.60 + j * 0.46, cw - 0.22, 0.38,
                   size=12, color=OFF_WHITE)


# ══════════════════════════════════════════════════════════════════════════════
#  SLIDE 6 — AI / RAG ENGINE
# ══════════════════════════════════════════════════════════════════════════════
def slide_6(prs):
    sl = prs.slides.add_slide(blank_layout(prs))
    set_slide_bg(sl, NAVY)
    add_slide_header(sl, 6, "AI-POWERED CONTRACT ANALYSIS",
                     "The RAG + Rule Engine pipeline that drives risk detection")
    add_footer(sl)

    components = [
        ("NLP Processing",      "Tokenizes & segments raw contract text",              INDIGO_LT),
        ("Embeddings",          "Converts clause text into high-dimensional vectors",  RGBColor(0x0E, 0x7A, 0xCA)),
        ("ChromaDB",            "Persistent vector store — semantic clause retrieval", RGBColor(0x6D, 0x28, 0xD9)),
        ("RAG Context Layer",   "Retrieves top-3 similar legal-risk knowledge chunks", PURPLE_DIM),
        ("LLM (GPT-4o-mini)",   "Explains risk & generates redline suggestion",        RED_HIGH),
        ("Rule Engine (40%)",   "Deterministic regex patterns for known risk terms",   AMBER_MED),
    ]
    for i, (name, desc, col) in enumerate(components):
        y = 1.35 + i * 1.0
        add_rect(sl, 0.35, y + 0.08, 0.10, 0.60, fill=col)
        add_tb(sl, name, 0.60, y + 0.05, 6.0, 0.38,
               size=14, bold=True, color=col)
        add_tb(sl, desc, 0.60, y + 0.46, 6.0, 0.40,
               size=11, color=GRAY_LIGHT, italic=True)

    # RAG pipeline visual
    add_tb(sl, "RAG PIPELINE", 7.5, 1.28, 5.5, 0.32,
           size=12, bold=True, color=GRAY_MED, align=PP_ALIGN.CENTER)

    rag = [
        ("CLAUSE TEXT",              SLATE),
        ("EMBEDDING GENERATION",     INDIGO),
        ("CHROMADB VECTOR STORE",    RGBColor(0x6D, 0x28, 0xD9)),
        ("SIMILAR CLAUSE RETRIEVAL", PURPLE_DIM),
        ("LLM + CONTEXT PROMPT",     CRIMSON),
        ("RISK EXPLANATION",         TEAL_DARK),
    ]
    for i, (lbl, col) in enumerate(rag):
        y = 1.72 + i * 0.90
        add_rect(sl, 7.42, y, 5.55, 0.62, fill=col)
        add_tb(sl, lbl, 7.42, y + 0.12, 5.55, 0.38,
               size=12, bold=True, color=WHITE, align=PP_ALIGN.CENTER)
        if i < 5:
            add_tb(sl, "▼", 10.12, y + 0.64, 0.55, 0.28,
                   size=12, color=INDIGO_LT, align=PP_ALIGN.CENTER)

    add_tb(sl, "Composite Score =  Rule Engine × 0.40  +  LLM Score × 0.40  +  RAG Score × 0.20",
           0.35, 6.88, 12.63, 0.28, size=10, color=GRAY_MED,
           align=PP_ALIGN.CENTER, italic=True)


# ══════════════════════════════════════════════════════════════════════════════
#  SLIDE 7 — RISK SCORING & REDLINE
# ══════════════════════════════════════════════════════════════════════════════
def slide_7(prs):
    sl = prs.slides.add_slide(blank_layout(prs))
    set_slide_bg(sl, NAVY)
    add_slide_header(sl, 7, "RISK ANALYSIS",
                     "Composite scoring formula + AI-generated redline alternatives")
    add_footer(sl)

    # Score formula card
    add_rect(sl, 0.35, 1.32, 6.05, 2.55, fill=CHARCOAL, line=INDIGO, line_w=0.8)
    add_tb(sl, "COMPOSITE RISK SCORE FORMULA", 0.58, 1.46, 5.65, 0.36,
           size=13, bold=True, color=INDIGO_LT)
    add_mltb(sl, [
        {'t': "Overall Score =", 's': 13, 'c': WHITE},
        {'t': "  (Rule Engine  ×  0.40)", 's': 13, 'c': AMBER_MED},
        {'t': "+ (LLM Score    ×  0.40)", 's': 13, 'c': RED_HIGH},
        {'t': "+ (RAG Score    ×  0.20)", 's': 13, 'c': INDIGO_LT},
    ], 0.58, 1.88, 5.65, 1.05, def_size=13, font="Courier New")

    # Tier badges
    add_tb(sl, "RISK TIERS", 0.58, 3.0, 5.65, 0.30,
           size=12, bold=True, color=GRAY_MED)
    tiers = [("LOW\n0–30",    GREEN_OK, 0.35), ("MEDIUM\n31–65", AMBER_MED, 2.42), ("HIGH\n66–100", RED_HIGH, 4.49)]
    for lbl, col, x in tiers:
        add_rect(sl, x, 3.35, 1.88, 1.08, fill=col)
        add_tb(sl, lbl, x, 3.35, 1.88, 1.08,
               size=14, bold=True, color=WHITE, align=PP_ALIGN.CENTER)

    add_tb(sl, "15+ clause categories:  Indemnity · Liability · Termination · Non-Compete · IP · Data Protection · Force Majeure …",
           0.35, 4.6, 6.0, 0.42, size=10, color=GRAY_MED, italic=True)

    # Example redline card
    add_rect(sl, 6.7, 1.32, 6.28, 5.45, fill=CHARCOAL, line=PURPLE_DIM, line_w=0.5)
    add_tb(sl, "EXAMPLE REDLINE  (from redline_service.py)", 6.92, 1.44, 5.88, 0.34,
           size=11, bold=True, color=GRAY_MED)

    add_pill_label(sl, "ORIGINAL CLAUSE", 6.92, 1.86, 2.20, 0.28, bg=RED_HIGH, size=9)
    add_tb(sl, '"Each party shall indemnify and hold harmless the other against any and all losses, claims, damages, liabilities and expenses…"',
           6.92, 2.20, 5.88, 0.90, size=11, color=GRAY_LIGHT, italic=True)

    add_rect(sl, 6.92, 3.16, 5.88, 0.03, fill=PURPLE_DIM)

    add_pill_label(sl, "RISK IDENTIFIED", 6.92, 3.26, 2.00, 0.28, bg=AMBER_MED, size=9)
    add_tb(sl, "Unlimited mutual indemnity — potentially high financial exposure with no liability cap.",
           6.92, 3.60, 5.88, 0.58, size=11, color=OFF_WHITE)

    add_rect(sl, 6.92, 4.25, 5.88, 0.03, fill=PURPLE_DIM)

    add_pill_label(sl, "AI-SUGGESTED REDLINE", 6.92, 4.35, 2.55, 0.28, bg=GREEN_OK, size=9)
    add_tb(sl, '"Each party (Indemnifying Party) shall defend and indemnify the other against direct third-party losses arising directly from gross negligence or material breach, subject to the Limitation of Liability cap set forth herein."',
           6.92, 4.70, 5.88, 1.68, size=11, color=OFF_WHITE, italic=True)


# ══════════════════════════════════════════════════════════════════════════════
#  SLIDE 8 — KEY FEATURES / SCREENSHOTS
# ══════════════════════════════════════════════════════════════════════════════
def slide_8(prs):
    sl = prs.slides.add_slide(blank_layout(prs))
    set_slide_bg(sl, NAVY)
    add_slide_header(sl, 8, "KEY FEATURES",
                     "Actual running application — screenshots from the live system")
    add_footer(sl)

    screenshots = [
        (DASHBOARD_IMG,  "Risk Dashboard",      "Metrics · Risk Overview · Upload Zone"),
        (UPLOAD_IMG,     "Contract Upload",     "PDF / DOCX Ingestion & Analysis Pipeline"),
        (AGREEMENTS_IMG, "Legal Agreements",    "Contract Registry with Risk-Level Filtering"),
        (COMPARE_IMG,    "Contract Comparison", "Side-by-Side Risk Posture Comparison"),
    ]
    positions = [
        (0.30, 1.32), (6.90, 1.32),
        (0.30, 4.28), (6.90, 4.28),
    ]
    cw, ch = 6.45, 2.76

    for (img_path, label, sublabel), (x, y) in zip(screenshots, positions):
        # Card border
        add_rect(sl, x, y, cw, ch, fill=RGBColor(0x0A, 0x14, 0x22),
                 line=PURPLE_DIM, line_w=0.5)
        # Screenshot
        if os.path.exists(img_path):
            try:
                sl.shapes.add_picture(
                    img_path,
                    Inches(x + 0.05), Inches(y + 0.05),
                    Inches(cw - 0.10), Inches(ch - 0.62)
                )
            except Exception as e:
                print(f"  [WARN] {img_path}: {e}")
        # Label strip
        add_rect(sl, x, y + ch - 0.58, cw, 0.58, fill=CHARCOAL)
        add_tb(sl, label, x + 0.15, y + ch - 0.54, cw - 0.25, 0.30,
               size=13, bold=True, color=INDIGO_LT)
        add_tb(sl, sublabel, x + 0.15, y + ch - 0.24, cw - 0.25, 0.22,
               size=9, color=GRAY_MED, italic=True)


# ══════════════════════════════════════════════════════════════════════════════
#  SLIDE 9 — RESULTS & BENEFITS
# ══════════════════════════════════════════════════════════════════════════════
def slide_9(prs):
    sl = prs.slides.add_slide(blank_layout(prs))
    set_slide_bg(sl, NAVY)
    add_slide_header(sl, 9, "RESULTS & BENEFITS",
                     "What the system delivers")
    add_footer(sl)

    benefits = [
        "Automated first-pass contract review",
        "15+ legal clause categories detected automatically",
        "AI + Rule-based risk scoring (composite 0–100)",
        "Plain-English risk explanations for each clause",
        "AI-suggested safer clause alternatives (redlines)",
        "Structured PDF executive summary report",
        "Centralized dashboard for contract management",
    ]
    add_tb(sl, "SYSTEM CAPABILITIES", 0.35, 1.28, 6.5, 0.32,
           size=11, bold=True, color=GRAY_MED)
    for i, text in enumerate(benefits):
        y = 1.70 + i * 0.75
        add_rect(sl, 0.35, y + 0.03, 0.40, 0.40, fill=GREEN_OK)
        add_tb(sl, "✓", 0.35, y + 0.03, 0.40, 0.40,
               size=14, bold=True, color=WHITE, align=PP_ALIGN.CENTER)
        add_tb(sl, text, 0.88, y + 0.05, 5.65, 0.40,
               size=13, color=OFF_WHITE)

    # BEFORE card
    add_tb(sl, "BEFORE  vs  AFTER", 7.5, 1.28, 5.55, 0.32,
           size=12, bold=True, color=GRAY_MED, align=PP_ALIGN.CENTER)

    add_rect(sl, 7.0, 1.70, 6.0, 2.36, fill=RGBColor(0x5A, 0x1A, 0x1A))
    add_tb(sl, "BEFORE ContractAI", 7.18, 1.80, 5.62, 0.36,
           size=13, bold=True, color=RED_HIGH)
    for i, pt in enumerate(["📄  Long contract received",
                             "👁   Manual clause-by-clause review",
                             "⏱   Hours of legal work required",
                             "❌  Risk clauses potentially missed"]):
        add_tb(sl, pt, 7.18, 2.24 + i * 0.42, 5.62, 0.36,
               size=11, color=OFF_WHITE)

    add_tb(sl, "▼", 9.72, 4.10, 0.55, 0.35,
           size=20, color=GRAY_MED, align=PP_ALIGN.CENTER)

    # AFTER card
    add_rect(sl, 7.0, 4.48, 6.0, 2.36, fill=TEAL_DARK)
    add_tb(sl, "AFTER ContractAI", 7.18, 4.58, 5.62, 0.36,
           size=13, bold=True, color=GREEN_OK)
    for i, pt in enumerate(["📤  Upload PDF or DOCX file",
                             "🤖  AI analyzes in seconds",
                             "📊  Risk dashboard opened instantly",
                             "✅  Actionable redline insights ready"]):
        add_tb(sl, pt, 7.18, 5.02 + i * 0.42, 5.62, 0.36,
               size=11, color=OFF_WHITE)


# ══════════════════════════════════════════════════════════════════════════════
#  SLIDE 10 — CONCLUSION & FUTURE SCOPE
# ══════════════════════════════════════════════════════════════════════════════
def slide_10(prs):
    sl = prs.slides.add_slide(blank_layout(prs))
    set_slide_bg(sl, NAVY)
    add_slide_header(sl, 10, "CONCLUSION & FUTURE SCOPE",
                     "Summary and roadmap for ContractAI")

    # Conclusion card
    add_rect(sl, 0.35, 1.28, 12.63, 1.62, fill=CHARCOAL, line=INDIGO, line_w=0.8)
    add_rect(sl, 0.35, 1.28, 0.10, 1.62, fill=INDIGO_LT)
    add_tb(sl, "CONCLUSION", 0.62, 1.34, 12.2, 0.30, size=11, bold=True, color=INDIGO_LT)
    add_tb(sl,
           "ContractAI demonstrates how AI, NLP, RAG and rule-based analysis can be combined to perform automated "
           "preliminary contract risk analysis through a web-based platform. The system integrates FastAPI, ChromaDB, "
           "OpenAI GPT-4o-mini, React/TypeScript and ReportLab to deliver risk scoring, clause detection, "
           "plain-English explanations, AI redlines and PDF reports.",
           0.62, 1.65, 12.2, 1.10, size=12, color=OFF_WHITE)

    # Future scope
    add_tb(sl, "FUTURE SCOPE", 0.35, 3.08, 12.63, 0.32,
           size=14, bold=True, color=INDIGO_LT)

    future_l = [
        "OCR support for scanned / image contracts",
        "Richer domain-specific legal knowledge base",
        "Fine-tuned legal-domain LLM models",
        "Multilingual contract analysis",
    ]
    future_r = [
        "Advanced contract comparison engine",
        "Human-in-the-loop legal review workflow",
        "Enterprise integrations (CRM / ERP / e-Sign)",
        "Role-based access control for legal teams",
    ]

    for i, item in enumerate(future_l):
        y = 3.52 + i * 0.63
        add_rect(sl, 0.35, y + 0.04, 0.36, 0.36, fill=INDIGO)
        add_tb(sl, "→", 0.35, y + 0.04, 0.36, 0.36,
               size=11, bold=True, color=WHITE, align=PP_ALIGN.CENTER)
        add_tb(sl, item, 0.82, y + 0.05, 5.7, 0.38, size=12, color=OFF_WHITE)

    for i, item in enumerate(future_r):
        y = 3.52 + i * 0.63
        add_rect(sl, 6.72, y + 0.04, 0.36, 0.36, fill=INDIGO)
        add_tb(sl, "→", 6.72, y + 0.04, 0.36, 0.36,
               size=11, bold=True, color=WHITE, align=PP_ALIGN.CENTER)
        add_tb(sl, item, 7.20, y + 0.05, 5.75, 0.38, size=12, color=OFF_WHITE)

    # Thank You section
    add_rect(sl, 0, 5.98, 13.33, 1.52, fill=DARK_PURPLE)
    add_tb(sl, "THANK YOU", 0, 6.06, 13.33, 0.60,
           size=30, bold=True, color=WHITE, align=PP_ALIGN.CENTER)
    add_tb(sl,
           "Questions?   ·   Ayush Yadav   ·   Roll No. 2329304   ·   B.Tech CSE 4-C   ·   Submitted to Ms. Mandeep Kaur",
           0, 6.65, 13.33, 0.30, size=11, color=GRAY_MED, italic=True, align=PP_ALIGN.CENTER)


# ══════════════════════════════════════════════════════════════════════════════
#  MAIN
# ══════════════════════════════════════════════════════════════════════════════
def main():
    print("=" * 60)
    print(" ContractAI — Professional PPTX Generator")
    print("=" * 60)
    prs = new_prs()

    for num, fn in [
        (1,  slide_1),
        (2,  slide_2),
        (3,  slide_3),
        (4,  slide_4),
        (5,  slide_5),
        (6,  slide_6),
        (7,  slide_7),
        (8,  slide_8),
        (9,  slide_9),
        (10, slide_10),
    ]:
        print(f"  Building slide {num} ...")
        fn(prs)

    prs.save(OUTPUT_PATH)
    print(f"\n✅ SUCCESS — {len(prs.slides)} slides saved to:")
    print(f"   {OUTPUT_PATH}")

if __name__ == "__main__":
    main()
